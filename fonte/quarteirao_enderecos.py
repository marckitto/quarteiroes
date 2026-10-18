# -*- coding: utf-8 -*-
"""
QUARTEIRÃO POR ENDEREÇO — Divisão de Controle de Vetores / Paulínia
====================================================================
Adiciona a cada endereço de uma planilha o número do quarteirão (e a zona),
usando:
  1) a malha de quarteirões da prefeitura (KML ou KMZ);
  2) o CNEFE 2022 do IBGE (endereços com coordenadas), que liga
     rua + número a um ponto no mapa;
  3) opcionalmente, uma base própria "endereço x quarteirão" já conferida.

TUDO roda no seu computador. Nenhum endereço é enviado para a internet.
(A única conexão possível é o download do CNEFE público do IBGE, se o
arquivo ainda não estiver na pasta.)

USO RÁPIDO (Prompt de Comando, na pasta do script):
    python quarteirao_enderecos.py enderecos.xlsx
ou arraste a planilha para cima do arquivo RODAR.bat.

Arquivos esperados na MESMA pasta do script (ou informe com --kml/--cnefe):
    - a malha de quarteirões (.kmz ou .kml)
    - 3536505_PAULINIA.zip ou 3536505_PAULINIA.csv (CNEFE) — baixado automaticamente se faltar
    - base_propria.xlsx (opcional) — colunas: RUA, NUMERO, QUARTEIRAO

Dependências (uma vez só):
    pip install pandas openpyxl shapely rapidfuzz
"""

import argparse
import io
import math
import os
import re
import sys
import unicodedata
import urllib.request
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter, defaultdict

try:
    import pandas as pd
    from rapidfuzz import fuzz, process
    from shapely.geometry import Point, Polygon, MultiPolygon
    from shapely.strtree import STRtree
except ImportError as e:  # pragma: no cover
    print("\nFalta instalar bibliotecas. Rode no Prompt de Comando:\n")
    print("    pip install pandas openpyxl shapely rapidfuzz\n")
    print(f"(detalhe: {e})")
    sys.exit(1)


# ----------------------------------------------------------------------------
# CONFIGURAÇÕES (ajuste se precisar)
# ----------------------------------------------------------------------------
CNEFE_URL = ("https://ftp.ibge.gov.br/Cadastro_Nacional_de_Enderecos_para_Fins_Estatisticos/"
             "Censo_Demografico_2022/Arquivos_CNEFE/CSV/Municipio/35_SP/3536505_PAULINIA.zip")
CNEFE_ARQUIVO = "3536505_PAULINIA.zip"

DIST_MAX_BORDA_M = 25      # ponto do CNEFE caído na rua: aceita o quarteirão mais próximo até X m
SIMILARIDADE_MIN = 87      # 0-100: abaixo disso a rua não é considerada encontrada
DIF_NUM_VIZINHO = 30       # diferença máxima de numeração para usar o "vizinho" mais próximo
CEPS_GENERICOS = {"13140000"}  # CEP geral da cidade: não serve para desempatar

# Palavras que aparecem nos nomes de colunas, para detecção automática
COLS_AUTO = {
    "rua": ["LOGRAD", "RUA", "ENDERE", "LOGR"],
    "numero": ["NUMERO", "NUM", "NU_", "NRO", "Nº"],
    "bairro": ["BAIRRO"],
    "cep": ["CEP"],
}


# ----------------------------------------------------------------------------
# LIMPEZA DE TEXTO
# ----------------------------------------------------------------------------
# Acentos que viraram símbolos (planilhas exportadas em codificação DOS/850)
_MOJIBAKE = {'┴': 'Á', '└': 'À', '┬': 'Â', '├': 'Ã', '╔': 'É', '╩': 'Ê', '═': 'Í',
             'Ë': 'Ó', 'È': 'Ô', 'ı': 'Õ', '┌': 'Ú', '▄': 'Ü', 'ß': 'Á', 'Ò': 'Ã',
             'þ': 'Ç', 'Û': 'Ê', 'Ý': 'Í', '¾': 'Ó', '¶': 'Ô', '§': 'Õ', '·': 'Ú', '³': 'Ü'}

TIPOS_LOGRADOURO = {
    "RUA": "RUA", "R": "RUA", "RU": "RUA",
    "AVENIDA": "AVENIDA", "AV": "AVENIDA", "AVE": "AVENIDA", "AVN": "AVENIDA", "AVDA": "AVENIDA",
    "ALAMEDA": "ALAMEDA", "AL": "ALAMEDA",
    "TRAVESSA": "TRAVESSA", "TV": "TRAVESSA", "TRAV": "TRAVESSA",
    "PRACA": "PRACA", "PCA": "PRACA", "PC": "PRACA",
    "ESTRADA": "ESTRADA", "EST": "ESTRADA", "ESTR": "ESTRADA",
    "RODOVIA": "RODOVIA", "ROD": "RODOVIA",
    "VIELA": "VIELA", "VIA": "VIA", "LARGO": "LARGO", "LGO": "LARGO",
    "PASSAGEM": "PASSAGEM", "VILA": "VILA", "VL": "VILA", "CAMINHO": "CAMINHO",
}
TITULOS = {
    "DR", "DRA", "DOUTOR", "DOUTORA", "PROF", "PROFA", "PROFESSOR", "PROFESSORA",
    "PREF", "PREFEITO", "PREFEITA", "ENG", "ENGENHEIRO", "ENGENHEIRA", "PADRE", "PE",
    "CEL", "CORONEL", "GAL", "GEN", "GENERAL", "VEREADOR", "VEREADORA", "VER",
    "DEP", "DEPUTADO", "DEPUTADA", "SEN", "SENADOR", "PRES", "PRESIDENTE",
    "MAJ", "MAJOR", "CAP", "CAPITAO", "TEN", "TENENTE", "SGT", "SARGENTO",
    "MAL", "MARECHAL", "ALM", "ALMIRANTE", "COMEND", "COMENDADOR", "MONS", "MONSENHOR",
    "FREI", "IRMA", "IRMAO", "DOM", "DONA", "SR", "SRA", "MESTRE", "MAESTRO",
    "PASTOR", "BISPO", "CONEGO", "DESEMBARGADOR", "JUIZ", "MINISTRO", "GOVERNADOR",
    "BRIGADEIRO", "COMANDANTE", "SOLDADO", "CABO", "PROFESSORES", "ENFERMEIRA",
}
ABREVIACOES = {"STA": "SANTA", "STO": "SANTO", "S": "SAO", "N": "NOSSA", "SRA": "SENHORA",
               "JD": "JARDIM", "PQ": "PARQUE", "RES": "RESIDENCIAL", "VL": "VILA",
               "MTE": "MONTE", "CJ": "CONJUNTO", "COND": "CONDOMINIO"}
STOPWORDS = {"DE", "DA", "DO", "DOS", "DAS", "E", "D"}
ROMANOS = {"I": "1", "II": "2", "III": "3", "IV": "4", "V": "5", "VI": "6",
           "VII": "7", "VIII": "8", "IX": "9", "X": "10"}


def corrigir_codificacao(s: str) -> str:
    s = "".join(_MOJIBAKE.get(ch, ch) for ch in s)
    # "GONÃALVES" -> "GONÇALVES" (ÃA / ÃU não existem em português)
    s = re.sub(r"Ã(?=[AU])", "Ç", s)
    return s


def sem_acento(s: str) -> str:
    s = unicodedata.normalize("NFKD", s)
    return "".join(c for c in s if not unicodedata.combining(c))


def texto_basico(s) -> str:
    if s is None or (isinstance(s, float) and math.isnan(s)):
        return ""
    s = str(s).upper().strip()
    s = corrigir_codificacao(s)
    s = sem_acento(s)
    s = s.replace("'", "")            # D'OTTAVIANO -> DOTTAVIANO
    s = re.sub(r"[^A-Z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def fonetica(token: str) -> str:
    """Aproxima grafias: CAZELATTO = CASELATO, OSORIIO = OSORIO, YPE = IPE."""
    t = token
    if t.isdigit():
        return t
    t = t.replace("PH", "F").replace("TH", "T").replace("Y", "I").replace("W", "V").replace("K", "C")
    t = re.sub(r"Ç", "S", t)
    t = re.sub(r"C(?=[EI])", "S", t)
    t = t.replace("Z", "S")
    t = re.sub(r"^H", "", t)
    t = re.sub(r"(.)\1+", r"\1", t)   # letras dobradas
    return t


def chave_rua(nome: str):
    """Devolve (chave_normalizada, tipo_logradouro) de um nome de rua."""
    s = texto_basico(nome)
    toks = s.split()
    tipo = ""
    while len(toks) > 1 and toks[0] in TIPOS_LOGRADOURO:   # "RUA AV PRES ..." -> tira os dois
        tipo = tipo or TIPOS_LOGRADOURO[toks[0]]
        toks = toks[1:]
    # título digitado errado no começo ("PREOFESSOR", "DOUTO") também sai
    while len(toks) > 2 and any(fuzz.ratio(toks[0], t) >= 85 for t in TITULOS if len(t) >= 6):
        toks = toks[1:]
    out = []
    for t in toks:
        if t in TITULOS or t in STOPWORDS:
            continue
        t = ABREVIACOES.get(t, t) if t not in ("S", "N") else t
        t = ROMANOS.get(t, t)
        out.append(fonetica(t))
    # descarta "S" e "N" soltos que sobraram
    out = [t for t in out if t not in ("S", "N")] or out
    return " ".join(out), tipo


def limpar_numero(v) -> str:
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return ""
    s = str(v).strip().upper()
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    return s


def separar_rua_numero(rua_bruta: str, num_bruto: str):
    """
    Limpa a rua (telefone, número colado, complemento) e decide o número.
    Retorna (rua_limpa, numero_int_ou_None, sufixo_letra, obs)
    """
    obs = []
    rua = corrigir_codificacao(str(rua_bruta or "").upper())
    # remove telefones: (19) 9 9123-8598, 99123-8598, 3874-1234
    tel = re.compile(r"\(?\b\d{2}\)?\s*9?\s*\d{4}\s*-\s*\d{4}\b|\b9?\d{4}-\d{4}\b")
    if tel.search(rua):
        rua = tel.sub(" ", rua)
        obs.append("telefone removido da rua")
    rua = re.sub(r"\(\s*\)", " ", rua)

    num_rua = None
    # número depois de vírgula: "RUA X, 41" / "AV Y , 2689 BETEL CASA E 29"
    m = re.search(r",\s*(?:N[º°O]?\.?\s*)?(\d{1,5})\b(.*)$", rua)
    if m:
        num_rua = m.group(1)
        if m.group(2).strip():
            obs.append(f"complemento ignorado: '{m.group(2).strip()}'")
        rua = rua[:m.start()]
    else:
        # número no fim: "R IDA PIETROBOM PIVA 48", "BORDIN140", "N 120"
        m = re.search(r"(?:\s|(?<=[A-Z]))(?:N[º°O]?\.?\s*)?(\d{1,5})\s*$", rua)
        if m and not re.search(r"\b(BRASIL|ESTRADA|RODOVIA|SP|KM)\s*\d*\s*$", rua[:m.end()]):
            num_rua = m.group(1)
            rua = rua[:m.start()]
    rua = rua.strip(" ,;-.:")
    rua = re.sub(r"^(RUA|AV|AVENIDA|R)\s*[:.]\s*", r"\1 ", rua)   # "RUA:" / "AV:" / "AV."

    num = limpar_numero(num_bruto)
    sufixo = ""
    n_int = None
    m2 = re.fullmatch(r"(\d+)\s*([A-Z]?)", num)
    if m2 and int(m2.group(1)) > 0:
        n_int = int(m2.group(1))
        sufixo = m2.group(2)
    if n_int is None and num_rua:
        n_int = int(num_rua)
        obs.append("número tirado do campo rua")
    elif n_int is not None and num_rua and int(num_rua) != n_int:
        obs.append(f"rua tinha outro número ({num_rua}); usado {n_int}")
    if n_int is None:
        obs.append("sem número")
    return rua, n_int, sufixo, obs


# ----------------------------------------------------------------------------
# MALHA DE QUARTEIRÕES (KML/KMZ)
# ----------------------------------------------------------------------------
LAT0 = -22.76  # Paulínia: projeção local em metros (suficiente para a cidade)
_KX = 111320.0 * math.cos(math.radians(LAT0))
_KY = 110540.0


def para_metros(lon, lat):
    return (lon * _KX, lat * _KY)


def _coords(txt):
    pts = []
    for c in txt.split():
        p = c.split(",")
        if len(p) >= 2:
            pts.append(para_metros(float(p[0]), float(p[1])))
    return pts


def ler_malha(caminho):
    if caminho.lower().endswith(".kmz"):
        with zipfile.ZipFile(caminho) as z:
            nome = next(n for n in z.namelist() if n.lower().endswith(".kml"))
            dados = z.read(nome)
    else:
        with open(caminho, "rb") as f:
            dados = f.read()
    raiz = ET.fromstring(dados)
    ns = {"k": raiz.tag.split("}")[0].strip("{")} if raiz.tag.startswith("{") else {"k": ""}
    q = "k:" if ns["k"] else ""
    geoms, attrs = [], []
    for pm in raiz.iter(f"{{{ns['k']}}}Placemark" if ns["k"] else "Placemark"):
        campos = {}
        for sd in pm.iter(f"{{{ns['k']}}}SimpleData" if ns["k"] else "SimpleData"):
            campos[(sd.get("name") or "").upper()] = (sd.text or "").strip()
        for d in pm.iter(f"{{{ns['k']}}}Data" if ns["k"] else "Data"):
            v = d.find(f"{q}value", ns)
            campos[(d.get("name") or "").upper()] = (v.text or "").strip() if v is not None else ""
        nome = pm.find(f"{q}name", ns)
        cod = next((v for k, v in campos.items() if k.startswith("QUARTEIR")), None)
        if cod is None:
            cod = (nome.text or "").strip() if nome is not None else ""
        zona = campos.get("ZONA", "")
        polys = []
        for pg in pm.iter(f"{{{ns['k']}}}Polygon" if ns["k"] else "Polygon"):
            ext = pg.find(f"{q}outerBoundaryIs/{q}LinearRing/{q}coordinates", ns)
            if ext is None:
                continue
            furos = [_coords(i.text) for i in pg.findall(f"{q}innerBoundaryIs/{q}LinearRing/{q}coordinates", ns)]
            p = Polygon(_coords(ext.text), [f for f in furos if len(f) >= 3])
            if not p.is_valid:
                p = p.buffer(0)
            polys.append(p)
        if not polys:
            continue
        g = polys[0] if len(polys) == 1 else MultiPolygon(
            [pp for p in polys for pp in (p.geoms if hasattr(p, "geoms") else [p])])
        geoms.append(g)
        attrs.append((cod, zona))
    return geoms, attrs


class Malha:
    def __init__(self, caminho):
        self.geoms, self.attrs = ler_malha(caminho)
        self.arvore = STRtree(self.geoms)
        self.sem_numero = {i for i, (c, _) in enumerate(self.attrs) if not re.search(r"\d", c)}

    def candidatos(self, lon, lat):
        """
        Lista [(quarteirao, distancia_m), ...] do mais próximo ao mais distante.
        Ponto dentro de um quarteirão -> [(q, 0.0)].
        Ponto na rua (fora dos polígonos) -> todos os quarteirões até DIST_MAX_BORDA_M.
        """
        pt = Point(para_metros(lon, lat))
        dentro = [i for i in self.arvore.query(pt, predicate="intersects")]
        dentro_validos = [i for i in dentro if i not in self.sem_numero]
        if dentro_validos:
            return [(self.attrs[dentro_validos[0]][0], 0.0)]
        if dentro:   # caiu num polígono "(XXXX)" — área sem número
            return []
        melhores = {}
        for i in self.arvore.query(pt.buffer(DIST_MAX_BORDA_M)):
            if i in self.sem_numero:
                continue
            d = self.geoms[i].distance(pt)
            c = self.attrs[i][0]
            if d <= DIST_MAX_BORDA_M and d < melhores.get(c, 1e9):
                melhores[c] = d
        return sorted(((c, round(d, 1)) for c, d in melhores.items()), key=lambda x: x[1])

    def localizar(self, lon, lat):
        """Retorna (quarteirao, distancia_m) do mais próximo, ou (None, None)."""
        c = self.candidatos(lon, lat)
        return c[0] if c else (None, None)


# ----------------------------------------------------------------------------
# CNEFE
# ----------------------------------------------------------------------------
def obter_cnefe(caminho):
    if os.path.exists(caminho):
        return caminho
    # aceita também o CSV já extraído, com o mesmo nome
    alt = os.path.splitext(caminho)[0] + ".csv"
    if os.path.exists(alt):
        return alt
    print(f"CNEFE não encontrado em '{caminho}'. Baixando do IBGE (arquivo público)...")
    try:
        urllib.request.urlretrieve(CNEFE_URL, caminho)
    except Exception as e:
        print("\nNão consegui baixar. Baixe manualmente e coloque na pasta do script:")
        print("   " + CNEFE_URL)
        print(f"(erro: {e})")
        sys.exit(1)
    return caminho


def ler_cnefe(caminho):
    if caminho.lower().endswith(".zip"):
        with zipfile.ZipFile(caminho) as z:
            nome = next(n for n in z.namelist() if n.lower().endswith(".csv"))
            bruto = z.read(nome)
    else:
        with open(caminho, "rb") as f:
            bruto = f.read()
    for enc in ("utf-8", "latin-1"):
        try:
            txt = bruto.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sep = ";" if txt[:2000].count(";") > txt[:2000].count(",") else ","
    df = pd.read_csv(io.StringIO(txt), sep=sep, dtype=str, low_memory=False)
    df.columns = [c.strip().upper() for c in df.columns]
    return df


def montar_referencia(cnefe, malha, log):
    """Cruza cada ponto do CNEFE com a malha -> tabela rua+número -> quarteirão."""
    col = lambda *nomes: next((c for c in nomes if c in cnefe.columns), None)
    c_tipo, c_tit, c_nome = col("NOM_TIPO_SEGLOGR"), col("NOM_TITULO_SEGLOGR"), col("NOM_SEGLOGR")
    c_num, c_cep = col("NUM_ENDERECO"), col("CEP")
    c_lat, c_lon, c_nv = col("LATITUDE"), col("LONGITUDE"), col("NV_GEO_COORD")
    c_loc = col("DSC_LOCALIDADE")
    if not all([c_nome, c_num, c_lat, c_lon]):
        raise SystemExit(f"Colunas do CNEFE não reconhecidas: {list(cnefe.columns)[:30]}")

    df = cnefe
    if c_nv:
        df = df[df[c_nv].astype(str).str.strip().isin(["1", "2", "3", "4"])]
    df = df.dropna(subset=[c_lat, c_lon])
    log(f"CNEFE: {len(cnefe):,} registros; {len(df):,} com coordenada no nível do endereço/face")

    nome_completo = (df[c_tipo].fillna("") + " " if c_tipo else "") + \
                    (df[c_tit].fillna("") + " " if c_tit else "") + df[c_nome].fillna("")
    brutos = defaultdict(list)       # chave -> [(num, cep, lon, lat, candidatos)]
    nomes_exib = {}
    fora = 0
    cache = {}
    for nm, num, cep, lat, lon in zip(nome_completo, df[c_num], df[c_cep] if c_cep else [""] * len(df),
                                      df[c_lat], df[c_lon]):
        try:
            la, lo = float(str(lat).replace(",", ".")), float(str(lon).replace(",", "."))
        except ValueError:
            continue
        k_xy = (round(lo, 6), round(la, 6))
        if k_xy not in cache:
            cache[k_xy] = malha.candidatos(lo, la)
        cands = cache[k_xy]
        if not cands:
            fora += 1
            continue
        chave, _ = chave_rua(nm)
        if not chave:
            continue
        try:
            n = int(float(num))
        except (TypeError, ValueError):
            n = 0
        brutos[chave].append((n, str(cep or "").replace("-", "")[:8], lo, la, cands))
        nomes_exib.setdefault(chave, re.sub(r"\s+", " ", nm).strip())

    # Pontos que caíram NA RUA (fora dos polígonos) podem estar mais perto do
    # quarteirão da frente. Resolve olhando os vizinhos do mesmo lado da rua
    # (mesma paridade) que caíram DENTRO de um quarteirão.
    ref = defaultdict(list)          # chave -> [(num, quarteirao, None, cep, lon, lat)]
    ajustados = 0
    for chave, regs in brutos.items():
        confiaveis = defaultdict(list)   # paridade -> [(num, quarteirao)]
        for n, _, _, _, cands in regs:
            if n > 0 and cands[0][1] == 0.0:
                confiaveis[n % 2].append((n, cands[0][0]))
        for lst in confiaveis.values():
            lst.sort()
        for n, cep, lo, la, cands in regs:
            q = cands[0][0]
            if cands[0][1] > 0 and len(cands) > 1 and n > 0 and confiaveis.get(n % 2):
                lst = confiaveis[n % 2]
                viz = sorted(lst, key=lambda x: abs(x[0] - n))[:4]
                blocos_viz = [b for _, b in viz]
                opcoes = [c for c, _ in cands]
                escolha = next((b for b in blocos_viz if b in opcoes), None)
                if escolha and escolha != q:
                    q = escolha
                    ajustados += 1
            ref[chave].append((n, q, None, cep, lo, la))
    log(f"Pontos sem quarteirão (fora da malha ou em área sem número): {fora:,}")
    log(f"Pontos na rua corrigidos pelo lado da rua: {ajustados:,}")
    log(f"Ruas na base de referência: {len(ref):,}")
    return ref, nomes_exib


def carregar_base_propria(caminho, ref, nomes_exib, malha_zona, log):
    """Base própria: colunas RUA, NUMERO, QUARTEIRAO (prioridade sobre o CNEFE)."""
    if not caminho or not os.path.exists(caminho):
        return {}
    b = pd.read_excel(caminho, dtype=str) if caminho.lower().endswith((".xlsx", ".xls")) \
        else pd.read_csv(caminho, dtype=str, sep=None, engine="python")
    b.columns = [c.strip().upper() for c in b.columns]
    cr = next(c for c in b.columns if "RUA" in c or "LOGRAD" in c)
    cn = next(c for c in b.columns if "NUM" in c)
    cq = next(c for c in b.columns if "QUART" in c)
    prop = defaultdict(dict)
    for r, n, q in zip(b[cr], b[cn], b[cq]):
        chave, _ = chave_rua(r)
        try:
            n = int(float(str(n)))
        except ValueError:
            continue
        q = str(q).strip()
        if q.isdigit():
            q = q.zfill(4)
        prop[chave][n] = q
        nomes_exib.setdefault(chave, str(r).strip().upper())
        ref.setdefault(chave, [])
    log(f"Base própria: {sum(len(v) for v in prop.values()):,} endereços")
    return prop


# ----------------------------------------------------------------------------
# CASAMENTO DOS ENDEREÇOS
# ----------------------------------------------------------------------------
class Casador:
    def __init__(self, ref, nomes_exib, propria, zona_por_qt, malha=None):
        self.malha = malha
        self.geom_por_qt = defaultdict(list)
        if malha:
            for g, (c, _) in zip(malha.geoms, malha.attrs):
                self.geom_por_qt[c].append(g)
        self.ref = ref
        self.nomes = nomes_exib
        self.prop = propria
        self.zona_por_qt = zona_por_qt
        self.chaves = list(ref.keys())
        self.chaves_sem_esp = [k.replace(" ", "") for k in self.chaves]
        self.ceps_por_chave = {k: Counter(r[3] for r in v) for k, v in ref.items()}
        self.cache = {}

    def achar_rua(self, chave, cep):
        if chave in self.ref:
            return chave, 100.0, ""
        if not chave:
            return None, 0.0, "rua vazia"
        if (chave, cep) in self.cache:
            return self.cache[(chave, cep)]
        # pré-seleção rápida por 3 critérios, depois pontuação fina
        sem_esp = chave.replace(" ", "")
        pre = set()
        for scorer, alvo, lista in ((fuzz.WRatio, chave, self.chaves),
                                    (fuzz.ratio, sem_esp, self.chaves_sem_esp),
                                    (fuzz.token_set_ratio, chave, self.chaves)):
            for _, _, idx in process.extract(alvo, lista, scorer=scorer, limit=8):
                pre.add(self.chaves[idx])
        q_toks = chave.split()
        pont = []
        for k in pre:
            k_toks = k.split()
            # iniciais: "JUSCELINO K OLIVEIRA" ~ "JUSCELINO CUBITSCHEC OLIVEIRA"
            q_exp = " ".join(next((kt for kt in k_toks if len(t) == 1 and kt.startswith(t)), t)
                             if len(t) == 1 else t for t in q_toks)
            s_sort = fuzz.token_sort_ratio(q_exp, k)
            s_junto = fuzz.ratio(q_exp.replace(" ", ""), k.replace(" ", ""))
            s = max(s_sort, s_junto)
            # nome incompleto ("IZAURA AMBROSIO" ~ "IZAURA AMBROSIO GONSALVES"):
            # só vale se o pedaço em comum tiver 2+ palavras e 10+ letras
            comuns = [t for t in q_toks if any(fuzz.ratio(t, kt) >= 85 for kt in k_toks)]
            if len(comuns) >= 2 and len("".join(comuns)) >= 10 and len(comuns) >= len(q_toks) - 1:
                s = max(s, fuzz.token_set_ratio(q_exp, k) - 6)
            if cep and cep not in CEPS_GENERICOS and self.ceps_por_chave[k].get(cep):
                s += 5   # mesmo CEP: pequeno bônus
            pont.append((min(round(s, 1), 99.9), k))
        pont.sort(reverse=True)
        res = (None, pont[0][0] if pont else 0.0, "rua não encontrada")
        if pont and pont[0][0] >= SIMILARIDADE_MIN:
            obs = ""
            if len(pont) > 1 and pont[1][0] >= SIMILARIDADE_MIN and pont[0][0] - pont[1][0] < 3:
                obs = f"rua ambígua (2ª opção: {self.nomes.get(pont[1][1])})"
            res = (pont[0][1], round(pont[0][0], 1), obs)
        self.cache[(chave, cep)] = res
        return res

    @staticmethod
    def _centro(regs, q):
        pts = [(r[4], r[5]) for r in regs if r[1] == q and r[4] is not None]
        if not pts:
            return None
        return (sum(p[0] for p in pts) / len(pts), sum(p[1] for p in pts) / len(pts))

    @staticmethod
    def _moda(regs):
        c = Counter(r[1] for r in regs)
        q, n = c.most_common(1)[0]
        return q, n / sum(c.values())

    def achar_quarteirao(self, chave, n, cep):
        """-> (quarteirao, metodo, confianca, obs, lon, lat)"""
        regs = self.ref.get(chave, [])
        prop = self.prop.get(chave, {})

        # 0) base própria, número exato
        if n is not None and n in prop:
            return prop[n], "base própria (número exato)", "ALTA", "", None, None

        # Se o CEP informado é específico, tenta primeiro só com os pontos desse CEP
        # (ajuda quando há ruas com o mesmo nome em lugares diferentes); se assim
        # o resultado ficar pior, usa a rua inteira.
        ordem = {"ALTA": 0, "MÉDIA": 1, "BAIXA": 2, "NÃO ENCONTRADO": 3}
        tentativas = [regs]
        if cep and cep not in CEPS_GENERICOS:
            filtrados = [r for r in regs if r[3] == cep]
            if filtrados and len(filtrados) < len(regs):
                tentativas = [filtrados, regs]
        melhor = None
        for t in tentativas:
            r = self._achar(t, prop, n)
            if melhor is None or ordem[r[2]] < ordem[melhor[2]]:
                melhor = r
            if ordem[r[2]] <= 1:
                break
        return melhor

    def _achar(self, regs_cep, prop, n):
        todos_q = Counter(r[1] for r in regs_cep)
        if n is None:
            if len(todos_q) == 1:
                q = next(iter(todos_q))
                return q, "rua toda num só quarteirão", "MÉDIA", "endereço sem número", None, None
            return None, "sem número", "NÃO ENCONTRADO", \
                f"rua passa por {len(todos_q)} quarteirões; precisa do número", None, None

        # 1) CNEFE, número exato
        exatos = [r for r in regs_cep if r[0] == n]
        if exatos:
            q, frac = self._moda(exatos)
            lon = sum(r[4] for r in exatos) / len(exatos)
            lat = sum(r[5] for r in exatos) / len(exatos)
            conf = "ALTA" if frac >= 0.75 else "MÉDIA"
            obs = "" if frac >= 0.75 else "pontos do mesmo número em quarteirões diferentes"
            return q, "CNEFE (número exato)", conf, obs, lon, lat

        # 2) regra da faixa: vizinhos do mesmo lado (mesma paridade)
        lado = {}
        for r in regs_cep:
            if r[0] > 0 and r[0] % 2 == n % 2:
                lado.setdefault(r[0], []).append(r)
        for num_p, q_p in prop.items():
            if num_p % 2 == n % 2:
                lado.setdefault(num_p, []).append((num_p, q_p, "", "", None, None))
        if not lado:
            return None, "sem vizinhos do mesmo lado", "NÃO ENCONTRADO", \
                "nenhum número do mesmo lado da rua na base", None, None
        nums = sorted(lado)
        abaixo = [x for x in nums if x < n]
        acima = [x for x in nums if x > n]
        qa = self._moda(lado[abaixo[-1]])[0] if abaixo else None
        qb = self._moda(lado[acima[0]])[0] if acima else None
        pa = self._centro(lado[abaixo[-1]], qa) if abaixo else None
        pb = self._centro(lado[acima[0]], qb) if acima else None
        lon = lat = None
        if pa and pb:   # posição estimada por interpolação da numeração
            t = (n - abaixo[-1]) / (acima[0] - abaixo[-1])
            lon, lat = pa[0] + t * (pb[0] - pa[0]), pa[1] + t * (pb[1] - pa[1])
        if qa and qb and qa == qb:
            return qa, f"faixa ({abaixo[-1]}–{acima[0]})", "ALTA", "", lon, lat
        # 3) vizinho mais próximo (mesmo lado da rua)
        cand = []
        if abaixo:
            cand.append((n - abaixo[-1], abaixo[-1], qa))
        if acima:
            cand.append((acima[0] - n, acima[0], qb))
        cand.sort()
        d, viz, q = cand[0]
        if d > DIF_NUM_VIZINHO:
            return q, f"vizinho distante (nº {viz})", "BAIXA", \
                f"número mais próximo na base está a {d} de distância", lon, lat
        if qa and qb:   # casa na troca de quarteirão
            return q, f"divisa (nº {viz})", "BAIXA", \
                f"entre o quarteirão {qa} (nº {abaixo[-1]}) e o {qb} (nº {acima[0]})", lon, lat
        # só existe vizinho de um lado (começo/fim da rua na base)
        conf = "MÉDIA"
        return q, f"vizinho mais próximo (nº {viz})", conf, "", lon, lat


def aprender(arquivo_resultado, caminho_base):
    """
    Lê a planilha de resultado (abas Enderecos e Revisar). Onde a coluna
    QUARTEIRAO_CORRIGIDO estiver preenchida, grava na base própria:
      - um número (ex.: 1075 ou 1075A) -> esse é o quarteirão certo;
      - "OK" -> confirma o quarteirão que o script sugeriu.
    Nas próximas execuções a base própria tem prioridade sobre o CNEFE.
    """
    novos = []
    for aba in ("Enderecos", "Revisar"):
        try:
            d = pd.read_excel(arquivo_resultado, sheet_name=aba, dtype=str)
        except ValueError:
            continue
        if "QUARTEIRAO_CORRIGIDO" not in d.columns:
            continue
        for _, r in d[d["QUARTEIRAO_CORRIGIDO"].fillna("").str.strip() != ""].iterrows():
            corr = str(r["QUARTEIRAO_CORRIGIDO"]).strip().upper()
            q = r.get("QUARTEIRAO") if corr == "OK" else corr
            if not isinstance(q, str) or not q.strip() or q.upper() == "NAN":
                continue
            q = q.strip().upper()
            if q.isdigit():
                q = q.zfill(4)
            rua = r.get("RUA_ENCONTRADA")
            if not isinstance(rua, str) or not rua.strip():
                rua = r.get("RUA_LIMPA")
            num = re.sub(r"\D", "", str(r.get("NUMERO_USADO") or ""))
            if not isinstance(rua, str) or not rua.strip() or not num:
                continue
            novos.append({"RUA": rua.strip(), "NUMERO": int(num), "QUARTEIRAO": q,
                          "FONTE": os.path.basename(arquivo_resultado),
                          "DATA": pd.Timestamp.now().strftime("%Y-%m-%d")})
    if not novos:
        print("Nenhuma linha com QUARTEIRAO_CORRIGIDO preenchido.")
        return
    novos = pd.DataFrame(novos)
    if os.path.exists(caminho_base):
        base = pd.read_excel(caminho_base, dtype=str)
        base = pd.concat([base, novos.astype(str)], ignore_index=True)
    else:
        base = novos.astype(str)
    base["_k"] = [chave_rua(r)[0] + "|" + str(n) for r, n in zip(base["RUA"], base["NUMERO"])]
    base = base.drop_duplicates("_k", keep="last").drop(columns="_k")
    base.to_excel(caminho_base, index=False)
    print(f"{len(novos)} correções gravadas. Base própria agora tem {len(base)} endereços: {caminho_base}")


# ----------------------------------------------------------------------------
# PRINCIPAL
# ----------------------------------------------------------------------------
def detectar_coluna(cols, tipo, informada):
    if informada:
        if informada not in cols:
            raise SystemExit(f"Coluna '{informada}' não existe. Colunas: {list(cols)}")
        return informada
    for c in cols:
        cu = sem_acento(str(c).upper())
        if any(p in cu for p in COLS_AUTO[tipo]):
            if tipo == "numero" and "CEP" in cu:
                continue
            if tipo == "rua" and ("NUM" in cu and "LOGR" not in cu):
                continue
            return c
    return None


def main():
    pasta = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser(description="Adiciona o quarteirão a uma lista de endereços.")
    ap.add_argument("enderecos", help="planilha .xlsx/.csv com os endereços")
    ap.add_argument("--kml", help="malha de quarteirões (.kmz/.kml)")
    ap.add_argument("--cnefe", default=os.path.join(pasta, CNEFE_ARQUIVO))
    ap.add_argument("--base", default=os.path.join(pasta, "base_propria.xlsx"))
    ap.add_argument("--aba", default=0, help="nome ou índice da aba (padrão: primeira)")
    ap.add_argument("--col-rua"); ap.add_argument("--col-numero")
    ap.add_argument("--col-bairro"); ap.add_argument("--col-cep")
    ap.add_argument("--saida")
    ap.add_argument("--aprender", action="store_true",
                    help="em vez de processar, lê a coluna QUARTEIRAO_CORRIGIDO da planilha "
                         "de resultado e grava as correções na base própria")
    a = ap.parse_args()
    if a.aprender:
        aprender(a.enderecos, a.base)
        return

    def log(msg):
        print(msg, flush=True)

    kml = a.kml or next((os.path.join(pasta, f) for f in sorted(os.listdir(pasta))
                         if f.lower().endswith((".kmz", ".kml"))), None)
    if not kml:
        raise SystemExit("Coloque a malha de quarteirões (.kmz/.kml) na pasta do script ou use --kml.")

    log(f"Malha: {os.path.basename(kml)}")
    malha = Malha(kml)
    log(f"  {len(malha.geoms):,} polígonos")
    zona_por_qt = {}
    for c, z in malha.attrs:
        zona_por_qt.setdefault(c, z)

    cnefe = ler_cnefe(obter_cnefe(a.cnefe))
    ref, nomes = montar_referencia(cnefe, malha, log)
    prop = carregar_base_propria(a.base, ref, nomes, zona_por_qt, log)
    casador = Casador(ref, nomes, prop, zona_por_qt, malha)

    aba = int(a.aba) if str(a.aba).isdigit() else a.aba
    if a.enderecos.lower().endswith(".csv"):
        df = pd.read_csv(a.enderecos, dtype=str, sep=None, engine="python")
    else:
        df = pd.read_excel(a.enderecos, sheet_name=aba, dtype=str)
    cr = detectar_coluna(df.columns, "rua", a.col_rua)
    cn = detectar_coluna(df.columns, "numero", a.col_numero)
    cb = detectar_coluna(df.columns, "bairro", a.col_bairro)
    cc = detectar_coluna(df.columns, "cep", a.col_cep)
    log(f"Colunas usadas -> rua: {cr} | número: {cn} | bairro: {cb} | CEP: {cc}")
    if not cr:
        raise SystemExit("Não achei a coluna da rua. Use --col-rua NOME_DA_COLUNA")

    saida = defaultdict(list)
    for _, linha in df.iterrows():
        rua, n, suf, obs = separar_rua_numero(linha[cr], linha[cn] if cn else "")
        cep = re.sub(r"\D", "", limpar_numero(linha[cc]) if cc else "")[:8]
        chave, _ = chave_rua(rua)
        k, sim, obs_rua = casador.achar_rua(chave, cep)
        if obs_rua:
            obs.append(obs_rua)
        if k is None:
            q, met, conf, o2, lon, lat = None, "rua não encontrada", "NÃO ENCONTRADO", "", None, None
        else:
            q, met, conf, o2, lon, lat = casador.achar_quarteirao(k, n, cep)
            if o2:
                obs.append(o2)
            if sim < 90 and conf in ("ALTA", "MÉDIA"):
                conf = "BAIXA"
                obs.append("nome da rua bem diferente do encontrado — conferir")
            elif conf == "ALTA" and sim < 95:
                conf = "MÉDIA"
                obs.append("nome da rua corrigido por semelhança — conferir")
        saida["QUARTEIRAO"].append(q)
        saida["QUARTEIRAO_CORRIGIDO"].append("")
        saida["ZONA"].append(zona_por_qt.get(q, "") if q else "")
        saida["CONFIANCA"].append(conf)
        saida["METODO"].append(met)
        saida["RUA_ENCONTRADA"].append(nomes.get(k, "") if k else "")
        saida["SIMILARIDADE_RUA"].append(sim if k else None)
        saida["NUMERO_USADO"].append(f"{n}{suf}" if n is not None else "")
        saida["RUA_LIMPA"].append(texto_basico(rua))
        saida["LATITUDE"].append(round(lat, 6) if lat else None)
        saida["LONGITUDE"].append(round(lon, 6) if lon else None)
        saida["OBS"].append("; ".join(o for o in obs if o))

    res = pd.concat([df.reset_index(drop=True), pd.DataFrame(saida)], axis=1)
    destino = a.saida or os.path.splitext(a.enderecos)[0] + "_com_quarteirao.xlsx"
    ordem = {"ALTA": 0, "MÉDIA": 1, "BAIXA": 2, "NÃO ENCONTRADO": 3}
    resumo = res["CONFIANCA"].value_counts().rename_axis("CONFIANCA").reset_index(name="QTD")
    resumo["%"] = (resumo["QTD"] / len(res) * 100).round(1)
    resumo = resumo.sort_values("CONFIANCA", key=lambda s: s.map(ordem))
    revisar = res[res["CONFIANCA"].isin(["BAIXA", "NÃO ENCONTRADO"]) |
                  res["OBS"].str.contains("conferir|ambígua", na=False)]
    with pd.ExcelWriter(destino, engine="openpyxl") as w:
        res.to_excel(w, sheet_name="Enderecos", index=False)
        revisar.to_excel(w, sheet_name="Revisar", index=False)
        resumo.to_excel(w, sheet_name="Resumo", index=False)
        from openpyxl.styles import PatternFill, Font
        cores = {"ALTA": "C6EFCE", "MÉDIA": "FFEB9C", "BAIXA": "F8CBAD", "NÃO ENCONTRADO": "D9D9D9"}
        for ws in w.book.worksheets:
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font = Font(bold=True)
                if c.value == "QUARTEIRAO_CORRIGIDO":
                    c.fill = PatternFill("solid", fgColor="BDD7EE")
            col_conf = next((c.column for c in ws[1] if c.value == "CONFIANCA"), None)
            if col_conf:
                for row in ws.iter_rows(min_row=2, min_col=col_conf, max_col=col_conf):
                    cor = cores.get(row[0].value)
                    if cor:
                        row[0].fill = PatternFill("solid", fgColor=cor)
            for col in ws.columns:
                larg = max(len(str(c.value or "")) for c in col[:200])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, larg + 2), 55)
    log("\nResultado:")
    for _, r in resumo.iterrows():
        log(f"  {r['CONFIANCA']:<15} {r['QTD']:>6}  ({r['%']}%)")
    log(f"\nArquivo gerado: {destino}")


if __name__ == "__main__":
    main()
